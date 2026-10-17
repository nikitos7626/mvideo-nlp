import os
import sys

if sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def generate_volgadev_deck(output_path="volgadev_checkpoint_15_10.pptx"):
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Цветовая палитра: светлый, чистый, высококонтрастный стиль VolgaDev / М.Видео
    COLOR_BG = RGBColor(248, 250, 252)          # Фон слайда #F8FAFC
    COLOR_CARD_BG = RGBColor(255, 255, 255)     # Фон карточки #FFFFFF
    COLOR_CARD_ALT = RGBColor(241, 245, 249)    # Акцентный серый #F1F5F9
    COLOR_DARK = RGBColor(15, 23, 42)           # Основной заголовок #0F172A
    COLOR_TEXT = RGBColor(51, 65, 85)           # Текст #334155
    COLOR_MUTED = RGBColor(100, 116, 139)       # Пояснения #64748B
    COLOR_RED = RGBColor(227, 6, 19)            # Фирменный красный М.Видео #E30613
    COLOR_BORDER = RGBColor(226, 232, 240)      # Рамки #E2E8F0
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_GREEN = RGBColor(16, 185, 129)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    def add_header_and_footer(slide, tag_text, title_text, slide_num_str):
        # Header Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_t = tag_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = tag_text.upper()
        p_t.font.name = "Arial"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_RED

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_DARK

        # Separator line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BORDER
        line.line.color.rgb = COLOR_BORDER

        # Footer Left
        f_left = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(5.0), Inches(0.35))
        p_fl = f_left.text_frame.paragraphs[0]
        p_fl.text = "/volgadev · checkpoint"
        p_fl.font.name = "Arial"
        p_fl.font.size = Pt(10)
        p_fl.font.bold = True
        p_fl.font.color.rgb = COLOR_MUTED

        # Footer Right
        f_right = slide.shapes.add_textbox(Inches(11.0), Inches(6.8), Inches(1.533), Inches(0.35))
        p_fr = f_right.text_frame.paragraphs[0]
        p_fr.text = slide_num_str
        p_fr.font.name = "Arial"
        p_fr.font.size = Pt(10)
        p_fr.font.bold = True
        p_fr.font.color.rgb = COLOR_MUTED
        p_fr.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # СЛАЙД 1: ТИТУЛЬНЫЙ (01 / 10)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(3.6), Inches(0.45))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = COLOR_RED
    badge1.line.fill.background()
    p_b1 = badge1.text_frame.paragraphs[0]
    p_b1.text = "CHECKPOINT 15.10 · VOLGADEV"
    p_b1.font.name = "Arial"
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = COLOR_WHITE
    p_b1.alignment = PP_ALIGN.CENTER

    tbox1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(2.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Заполним каталог М.Видео\nбез ручного труда"
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_DARK

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "\nИИ для автоматического заполнения характеристик товаров · заказчик М.Видео"
    p1_sub.font.name = "Arial"
    p1_sub.font.size = Pt(17)
    p1_sub.font.color.rgb = COLOR_TEXT

    # 3 акцентные карточки ценности
    s1_cards = [
        ("Точность в поиске", "Автоматически извлекаем обязательные атрибуты из текста карточек и выводим товары в фильтры каталога."),
        ("Канонизация значений", "Приводим разнобой «алый / red», «1024 Гб / 1 ТБ» к единому справочнику М.Видео без ручного труда."),
        ("Оценка уверенности", "Каждому атрибуту присваиваем оценку модели. Сомнительные карточки отправляем на ручную проверку.")
    ]
    for idx, (t, d) in enumerate(s1_cards):
        c = add_card(s1, 1.0 + idx * 3.85, 4.6, 3.6, 1.8)
        tf = c.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = t
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_RED
        pd = tf.add_paragraph()
        pd.text = d
        pd.font.size = Pt(11)
        pd.font.color.rgb = COLOR_TEXT

    # Footer slide 1
    f1_l = s1.shapes.add_textbox(Inches(1.0), Inches(6.8), Inches(5.0), Inches(0.35))
    f1_l.text_frame.paragraphs[0].text = "/volgadev · checkpoint"
    f1_l.text_frame.paragraphs[0].font.size = Pt(10)
    f1_l.text_frame.paragraphs[0].font.bold = True
    f1_l.text_frame.paragraphs[0].font.color.rgb = COLOR_MUTED

    f1_r = s1.shapes.add_textbox(Inches(10.8), Inches(6.8), Inches(1.5), Inches(0.35))
    f1_r.text_frame.paragraphs[0].text = "01 / 10"
    f1_r.text_frame.paragraphs[0].font.size = Pt(10)
    f1_r.text_frame.paragraphs[0].font.bold = True
    f1_r.text_frame.paragraphs[0].font.color.rgb = COLOR_MUTED
    f1_r.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # СЛАЙД 2: ПРОБЛЕМА (02 / 10)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header_and_footer(s2, "ПРОБЛЕМА", "Неровная карточка ведёт к потере продаж", "02 / 10")

    chain_items = [
        ("КАРТОЧКА", "Хаос в данных", "Разнобой «алый / красный / red», единицы вразнобой, важные поля спрятаны в описании."),
        ("АТРИБУТ НЕ ЗАПОЛНЕН", "Выпадение из поиска", "Товар не попадает в фильтры поиска каталога М.Видео."),
        ("ПОКУПАТЕЛЬ НЕ НАХОДИТ", "Падение конверсии", "Покупатель просто не видит товар в выдаче. Падают CTR и конверсия."),
        ("ПОТЕРЯ GMV", "Потеря выручки", "Теряем продажи маркетплейса и комиссию с каждого непроданного товара.")
    ]

    for idx, (step_t, step_sub, step_desc) in enumerate(chain_items):
        x = 0.8 + idx * 2.95
        c = add_card(s2, x, 2.0, 2.75, 4.3)
        tf = c.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"ЭТАП 0{idx+1}"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_RED

        p1 = tf.add_paragraph()
        p1.text = step_t
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_DARK

        p_sub = tf.add_paragraph()
        p_sub.text = step_sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_RED if idx == 3 else COLOR_TEXT

        p2 = tf.add_paragraph()
        p2.text = "\n" + step_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # СЛАЙД 3: РЕШЕНИЕ (03 / 10)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header_and_footer(s3, "РЕШЕНИЕ", "Из карточки — в аккуратные атрибуты", "03 / 10")

    # Слева: Сырой вход
    c3_in = add_card(s3, 0.8, 1.8, 5.7, 3.6, bg_color=COLOR_CARD_ALT)
    tf3_i = c3_in.text_frame
    tf3_i.word_wrap = True
    p0 = tf3_i.paragraphs[0]
    p0.text = "СЫРОЙ ВХОД"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED

    p_txt = tf3_i.add_paragraph()
    p_txt.text = "\n«Apple iPhone 15 Pro Max 256GB Black Titanium. Флагман, 6.7″ OLED retina, батарея 4422 мАч, цвет чёрный титан. rom 256gb, color black»"
    p_txt.font.size = Pt(13)
    p_txt.font.color.rgb = COLOR_DARK

    p_note = tf3_i.add_paragraph()
    p_note.text = "\nМаркетинговый текст, дубли в спецификациях продавца, синонимы цветов и латиница."
    p_note.font.size = Pt(11)
    p_note.font.color.rgb = COLOR_MUTED

    # Справа: Аккуратный JSON
    c3_out = add_card(s3, 6.8, 1.8, 5.75, 3.6)
    tf3_o = c3_out.text_frame
    tf3_o.word_wrap = True
    p0 = tf3_o.paragraphs[0]
    p0.text = "АККУРАТНЫЙ JSON"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_GREEN

    specs = [
        "• Бренд: Apple (уверенность 0.98)",
        "• Модель: iPhone 15 Pro Max (уверенность 0.91)",
        "• Объём памяти: 256 ГБ (уверенность 0.96)",
        "• Диагональ экрана: 6.7″ (уверенность 0.94)",
        "• Цвет: чёрный (уверенность 0.92, нормализовано)",
        "Каждому значению — откалиброванная уверенность модели."
    ]
    for s in specs:
        p = tf3_o.add_paragraph()
        p.text = s
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK

    # Нижняя плашка пайплайна
    c3_pipe = add_card(s3, 0.8, 5.65, 11.75, 0.95, bg_color=COLOR_CARD_ALT)
    tf_p = c3_pipe.text_frame
    tf_p.word_wrap = True
    p0 = tf_p.paragraphs[0]
    p0.text = "ПАЙПЛАЙН: Карточка → извлечение → нормализация → JSON с уверенностью."
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_DARK

    p1 = tf_p.add_paragraph()
    p1.text = "Низкая уверенность (< 0.75) или пропуск обязательного поля → отправка на ручную проверку."
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # СЛАЙД 4: ПОДХОД · ГИБРИД (04 / 10)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header_and_footer(s4, "ПОДХОД · ГИБРИД", "Гибрид, а не одна нейросеть", "04 / 10")

    cols4 = [
        ("ПРАВИЛА + СЛОВАРИ", "Дешёвое ядро каталога", "Бренды, цвета, размеры, единицы измерения. Быстро (доли миллисекунды) и без затрат на модель."),
        ("ZERO-SHOT NER", "Длинный хвост без дообучения", "Редкие значения и новые формулировки в свободном тексте — без затратного дообучения и риска оверфиттинга."),
        ("LLM — ТОЧЕЧНО", "Только для редких пропусков", "Только для пропусков и новых категорий, с жестким программным контролем галлюцинаций.")
    ]

    for idx, (t, sub, desc) in enumerate(cols4):
        x = 0.8 + idx * 4.0
        c = add_card(s4, x, 1.8, 3.75, 3.6)
        tf = c.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = t
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_RED

        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_DARK

        p_desc = tf.add_paragraph()
        p_desc.text = "\n" + desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_TEXT

    c4_bot = add_card(s4, 0.8, 5.65, 11.75, 0.95, bg_color=COLOR_CARD_ALT)
    tf4_b = c4_bot.text_frame
    tf4_b.word_wrap = True
    p0 = tf4_b.paragraphs[0]
    p0.text = "КЛЮЧЕВОЙ ПРИНЦИП"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED

    p1 = tf4_b.add_paragraph()
    p1.text = "Ядро дешёвое, дорогое — точечно. Это прямой ответ на ключевое требование кейса М.Видео «учитывайте стоимость на масштабе каталога»."
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # СЛАЙД 5: МЕТРИКИ · ПЛАН (05 / 10)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header_and_footer(s5, "МЕТРИКИ · ПЛАН", "Честная оценка и ясный график", "05 / 10")

    # Верхний блок метрик
    c5_met = add_card(s5, 0.8, 1.8, 11.75, 2.2)
    tf5_m = c5_met.text_frame
    tf5_m.word_wrap = True
    p0 = tf5_m.paragraphs[0]
    p0.text = "МЕТРИКИ (НЕ ОБЩАЯ ACCURACY)"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED

    met_text = [
        "• Per-attribute F1 на эталоне: целевой ориентир для ядра ≥ 0.85, для длинного хвоста ≥ 0.70.",
        "• Порог уверенности: калибруется на валидационном сете по балансу Precision / Recall.",
        "• Coverage (Покрытие): доля карточек с полностью заполненными обязательными полями.",
        "• Оценка по категориям: редкие сегменты и небольшие селлеры не обрабатываются хуже массовых."
    ]
    for m in met_text:
        p = tf5_m.add_paragraph()
        p.text = m
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK

    # Нижний блок таймлайна
    c5_time = add_card(s5, 0.8, 4.25, 11.75, 2.3, bg_color=COLOR_CARD_ALT)
    tf5_t = c5_time.text_frame
    tf5_t.word_wrap = True
    p0 = tf5_t.paragraphs[0]
    p0.text = "ГРАФИК ДО ФИНАЛА"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_DARK

    timeline_steps = [
        ("15.10", "Чекпоинт (сегодня)"),
        ("16.10", "Данные и поля от заказчика"),
        ("19.10", "Первый сквозной прогон"),
        ("23.10", "Лидерборд (оценим прогресс)"),
        ("31.10", "Финал, Магас (защита)")
    ]
    # Нарисуем таймлайн из 5 карточек
    for idx, (dt, event) in enumerate(timeline_steps):
        x = 1.0 + idx * 2.3
        c = add_card(s5, x, 4.8, 2.15, 1.5)
        tf = c.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = dt
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_RED
        pe = tf.add_paragraph()
        pe.text = event
        pe.font.size = Pt(10)
        pe.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # СЛАЙД 6: РИСКИ (06 / 10)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header_and_footer(s6, "РИСКИ", "Называем — и закрываем мерами", "06 / 10")

    risks_grid = [
        ("ГАЛЛЮЦИНАЦИИ LLM", "Значение должно строго встречаться в тексте карточки (span matching), иначе confidence = 0. Плюс жесткая схема вывода и порог уверенности."),
        ("СТОИМОСТЬ НА МАСШТАБЕ", "Гибрид: дешёвое ядро, дорогое — точечно. Правила работают со скоростью ~1000 карт/с, GLiNER ~50/с; LLM подключается только на ~5% пропусков."),
        ("ДЛИННЫЙ ХВОСТ", "Per-attribute F1 + gold-set отдельно по категориям вместо общей точности. Контроль редких наименований через канонизатор."),
        ("НОВЫЕ КАТЕГОРИИ", "Неочевидные и новые атрибуты оцениваем на gold-сете, размеченном вручную. Разметки возвращаются в петлю дообучения и правил.")
    ]

    for idx, (r_title, r_desc) in enumerate(risks_grid):
        row = idx // 2
        col = idx % 2
        x = 0.8 + col * 5.95
        y = 1.8 + row * 2.45
        c = add_card(s6, x, y, 5.75, 2.25)
        tf = c.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = r_title
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_RED

        p_desc = tf.add_paragraph()
        p_desc.text = "\n" + r_desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # СЛАЙД 7: КАК РАБОТАЕТ (07 / 10)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header_and_footer(s7, "КАК РАБОТАЕТ", "Открытый стек, дешёвое ядро — вперёд", "07 / 10")

    tech_cards = [
        ("ЯДРО — ПРАВИЛА", "FlashText + Regex", "Мгновенный поиск брендов и цветов по словарям за один проход O(N). Регулярные выражения для физических величин (ГБ, мАч, дюймы)."),
        ("ХВОСТ — ZERO-SHOT NER", "GLiNER Multi-v2.1", "Любые произвольные атрибуты без дообучения модели. Локальный инференс прямо на ВМ, возвращает математические вероятности спанов."),
        ("КАНОНИЗАЦИЯ", "RapidFuzz + Pydantic", "«сяоми» → «Xiaomi», «space gray» → «серый». UnitNormalizer: «1024 Гб» → «1 ТБ». Pydantic валидирует структуру и отсекает смысловые аномалии."),
        ("КОНТРОЛЬ", "LLM-фоллбэк", "LLM подключается только при реальных пропусках. Если ядро извлечено на первых шагах — дорогие модели не вызываются вовсе.")
    ]

    for idx, (t, tech, desc) in enumerate(tech_cards):
        row = idx // 2
        col = idx % 2
        x = 0.8 + col * 5.95
        y = 1.8 + row * 2.45
        c = add_card(s7, x, y, 5.75, 2.25)
        tf = c.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = t
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_RED

        p_t = tf.add_paragraph()
        p_t.text = tech
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_DARK

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # СЛАЙД 8: РЕЗУЛЬТАТ · ПОЧЕМУ МЫ (08 / 10)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header_and_footer(s8, "РЕЗУЛЬТАТ · ПОЧЕМУ МЫ", "Пример: iPhone 15 Pro Max", "08 / 10")

    # Слева: Таблица
    t_shape = s8.shapes.add_table(6, 4, Inches(0.8), Inches(1.8), Inches(6.2), Inches(4.7))
    tbl = t_shape.table

    tbl_headers = ["АТРИБУТ", "ЗНАЧЕНИЕ", "УВЕРЕННОСТЬ", "ИСТОЧНИК"]
    for c_idx, h in enumerate(tbl_headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER

    res_data = [
        ["Бренд", "Apple", "0.98", "FlashText (Rule)"],
        ["Модель", "iPhone 15 Pro Max", "0.91", "GLiNER"],
        ["Объём памяти", "256 ГБ", "0.96", "Regex + Normalizer"],
        ["Диагональ", "6.7″", "0.94", "Regex + Normalizer"],
        ["Цвет", "чёрный", "0.92", "RapidFuzz"]
    ]
    for r_idx, row in enumerate(res_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_DARK
            p.alignment = PP_ALIGN.CENTER

    # Справа: Карточка "Открытый стек + контроль"
    c8_r = add_card(s8, 7.3, 1.8, 5.25, 4.7)
    tf8_r = c8_r.text_frame
    tf8_r.word_wrap = True

    p0 = tf8_r.paragraphs[0]
    p0.text = "ОТКРЫТЫЙ СТЕК + КОНТРОЛЬ"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED

    r_points = [
        ("Zero-shot NER без дообучения:", "Быстрый старт, контроль стоимости и отсутствие привязки к платному API."),
        ("Анти-галлюцинатор:", "Значение обязано быть в тексте карточки, иначе уверенность принудительно сбрасывается в 0."),
        ("Порог 0.75 калибруется на валидации:", "В каталог на автомате идёт только надёжное. Сомнительное — на ручную проверку и в петлю дообучения."),
        ("Команда из трёх человек:", "Ядро, демо и продукт закрыты с первого дня хакатона.")
    ]
    for h, b in r_points:
        p = tf8_r.add_paragraph()
        p.text = "• " + h + " " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # СЛАЙД 9: КОМАНДА (09 / 10)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header_and_footer(s9, "КОМАНДА", "Кто за что отвечает", "09 / 10")

    team_roles = [
        ("NLP-ИНЖЕНЕР", "Ядро решения", [
            "Промпты и регулярки",
            "Извлечение атрибутов",
            "Нормализация величин",
            "Калибровка уверенности",
            "Генерация валидного JSON"
        ]),
        ("BACKEND-РАЗРАБОТЧИК", "Инфраструктура и API", [
            "API-сервис пайплайна",
            "Демо-интерфейс «карточка → JSON»",
            "Обработка ошибок и батчей",
            "Деплой на виртуальную машину",
            "Очередь модерации"
        ]),
        ("РМ / ПРОДУКТ", "Метрики и защита", [
            "Анализ данных и EDA",
            "Контроль метрик (F1, Coverage)",
            "Презентация и питч",
            "Контакт с заказчиком и жюри",
            "Соблюдение таймлайна"
        ])
    ]

    for idx, (role_t, role_sub, tasks) in enumerate(team_roles):
        x = 0.8 + idx * 4.0
        c = add_card(s9, x, 1.8, 3.75, 3.6)
        tf = c.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = role_t
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_RED

        p_sub = tf.add_paragraph()
        p_sub.text = role_sub
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_DARK

        for t in tasks:
            p = tf.add_paragraph()
            p.text = "• " + t
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT

    c9_bot = add_card(s9, 0.8, 5.65, 11.75, 0.95, bg_color=COLOR_CARD_ALT)
    tf9_b = c9_bot.text_frame
    tf9_b.word_wrap = True
    p0 = tf9_b.paragraphs[0]
    p0.text = "РЕШЕНИЯ БЫСТРЕЕ"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED

    p1 = tf9_b.add_paragraph()
    p1.text = "Команда из трёх человек — весь технологический стек на связи: ядро, демо и продукт закрыты с первого дня."
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_DARK

    # =========================================================================
    # СЛАЙД 10: ФИНАЛ (10 / 10)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)

    # Заголовок финала
    add_header_and_footer(s10, "ФИНАЛ", "Спасибо за внимание!", "10 / 10")

    c10 = add_card(s10, 1.0, 2.0, 11.333, 4.4, bg_color=COLOR_CARD_BG)
    tf10 = c10.text_frame
    tf10.word_wrap = True

    p0 = tf10.paragraphs[0]
    p0.text = "\nПонимание · Гибрид · Метрики · План · Команда"
    p0.font.size = Pt(28)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED
    p0.alignment = PP_ALIGN.CENTER

    p1 = tf10.add_paragraph()
    p1.text = "\n\nЧто дальше:"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_DARK
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf10.add_paragraph()
    p2.text = "16.10 — получаем данные и список полей от заказчика\n19.10 — первый сквозной прогон пайплайна «карточка → JSON»"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_TEXT
    p2.alignment = PP_ALIGN.CENTER

    prs.save(output_path)
    print(f"[+] Презентация чекпоинта успешно сгенерирована: {output_path}")


if __name__ == "__main__":
    generate_volgadev_deck()
